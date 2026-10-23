import os
import sys
import json
import base64
import tempfile
import threading
import webview
import fitz  # PyMuPDF
from http.server import SimpleHTTPRequestHandler
import socketserver

def convertir_a_pdf(ruta_origen):
    ext = os.path.splitext(ruta_origen)[1].lower()
    if ext == ".pdf":
        with open(ruta_origen, "rb") as f:
            return f.read()

    temp_dir = tempfile.gettempdir()
    nombre_base = os.path.splitext(os.path.basename(ruta_origen))[0]
    pdf_salida = os.path.join(temp_dir, f"temp_view_{nombre_base}.pdf")
    error_detalle = None

    if ext in [".ppt", ".pptx"] and sys.platform == "win32":
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()

            ppt_app = win32com.client.Dispatch("PowerPoint.Application")
            pres = ppt_app.Presentations.Open(
                os.path.abspath(ruta_origen),
                True, False, False
            )

            try:
                ppt_app.WindowState = 2
            except Exception:
                pass

            try:
                pres.SaveAs(os.path.abspath(pdf_salida), 32)
            except Exception:
                pres.ExportAsFixedFormat(os.path.abspath(pdf_salida), 2, 1)

            pres.Close()
            ppt_app.Quit()

            if os.path.exists(pdf_salida):
                with open(pdf_salida, "rb") as f:
                    data = f.read()
                try: os.remove(pdf_salida)
                except Exception: pass
                return data
        except Exception as e:
            error_detalle = f"PowerPoint: {e}"

    if sys.platform == "win32" and ext in [".doc", ".docx", ".xls", ".xlsx"]:
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()

            if ext in [".doc", ".docx"]:
                word = win32com.client.DispatchEx("Word.Application")
                word.Visible = False
                word.DisplayAlerts = 0
                doc = word.Documents.Open(os.path.abspath(ruta_origen), ReadOnly=True)
                doc.SaveAs(os.path.abspath(pdf_salida), FileFormat=17)
                doc.Close(False)
                word.Quit()

            elif ext in [".xls", ".xlsx"]:
                excel = win32com.client.DispatchEx("Excel.Application")
                excel.Visible = False
                excel.DisplayAlerts = False
                wb = excel.Workbooks.Open(os.path.abspath(ruta_origen), ReadOnly=True)
                wb.ExportAsFixedFormat(0, os.path.abspath(pdf_salida))
                wb.Close(False)
                excel.Quit()

            if os.path.exists(pdf_salida):
                with open(pdf_salida, "rb") as f:
                    data = f.read()
                try: os.remove(pdf_salida)
                except Exception: pass
                return data
        except Exception as e:
            error_detalle = f"Word/Excel: {e}"

    if ext == ".pptx":
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            prs = Presentation(ruta_origen)
            doc_pdf = fitz.open()

            for slide in prs.slides:
                pag = doc_pdf.new_page(width=1280, height=720)
                pag.draw_rect(fitz.Rect(0, 0, 1280, 720), fill=(1, 1, 1))

                for shape in slide.shapes:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        try:
                            img_bytes = shape.image.blob
                            x0 = (shape.left / prs.slide_width) * 1280
                            y0 = (shape.top / prs.slide_height) * 720
                            x1 = x0 + (shape.width / prs.slide_width) * 1280
                            y1 = y0 + (shape.height / prs.slide_height) * 720
                            pag.insert_image(fitz.Rect(x0, y0, x1, y1), stream=img_bytes)
                        except Exception:
                            pass

                    if shape.has_text_frame:
                        x0 = max(40, (shape.left / prs.slide_width) * 1280)
                        y0 = max(30, (shape.top / prs.slide_height) * 720)
                        w = min(1200, (shape.width / prs.slide_width) * 1280)
                        h = min(660, (shape.height / prs.slide_height) * 720)

                        texto_parrafos = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
                        if texto_parrafos:
                            pag.insert_textbox(
                                fitz.Rect(x0, y0, x0 + w, y0 + h),
                                "\n".join(texto_parrafos),
                                fontsize=16 if y0 < 120 else 13,
                                fontname="helv"
                            )

            if len(doc_pdf) > 0:
                return doc_pdf.tobytes()
        except Exception as e:
            error_detalle = f"PPTX fallback: {e}"

    if error_detalle:
        print(f"Error detallado: {error_detalle}")

    return None


class PDFViewerAPI:
    def __init__(self):
        self._window = None

    def set_window(self, window):
        self._window = window

    def forzar_cierre(self):
        if self._window:
            self._window.destroy()

    def abrir_dialogo_archivo(self):
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)

        tipos = [
            ("Todos los archivos compatibles", "*.pdf;*.docx;*.doc;*.xlsx;*.xls;*.pptx;*.ppt"),
            ("Presentaciones PowerPoint (*.pptx, *.ppt)", "*.pptx;*.ppt"),
            ("Documentos de Word (*.docx, *.doc)", "*.docx;*.doc"),
            ("Planillas de Excel (*.xlsx, *.xls)", "*.xlsx;*.xls"),
            ("Archivos PDF (*.pdf)", "*.pdf"),
            ("Todos los archivos", "*.*")
        ]

        rutas = filedialog.askopenfilenames(
            title="Seleccionar documentos o presentaciones",
            filetypes=tipos
        )
        root.destroy()

        if not rutas:
            return {"exito": False}

        resultados = []
        for ruta in rutas:
            if os.path.exists(ruta):
                bytes_pdf = convertir_a_pdf(ruta)
                if bytes_pdf:
                    b64 = base64.b64encode(bytes_pdf).decode("utf-8")
                    resultados.append({
                        "exito": True, 
                        "nombre": os.path.basename(ruta), 
                        "data": b64, 
                        "ruta": ruta
                    })
        
        if resultados:
            return {"exito": True, "archivos": resultados}
        
        return {"exito": False, "error": "No se pudo procesar ninguno de los archivos seleccionados"}

    def convertir_archivo(self, nombre, datos_b64, ruta=None):
        try:
            if ruta and os.path.exists(ruta):
                pdf_bytes = convertir_a_pdf(ruta)
                if pdf_bytes:
                    return {"exito": True, "data": base64.b64encode(pdf_bytes).decode("utf-8")}
            if datos_b64:
                ext = os.path.splitext(nombre)[1].lower()
                temp_in = os.path.join(tempfile.gettempdir(), f"input_conv_{nombre}")
                with open(temp_in, "wb") as f:
                    f.write(base64.b64decode(datos_b64))
                pdf_bytes = convertir_a_pdf(temp_in)
                try: os.remove(temp_in)
                except Exception: pass
                if pdf_bytes:
                    return {"exito": True, "data": base64.b64encode(pdf_bytes).decode("utf-8")}
        except Exception as e:
            pass
        return {"exito": False}

    def guardar_documento(self, ruta_doc, datos_b64, datos_json):
        if not ruta_doc or not datos_b64:
            return {"exito": False, "error": "No hay archivo cargado"}

        nombre_base = os.path.splitext(os.path.basename(ruta_doc))[0]
        resultado = self._window.create_file_dialog(
            webview.SAVE_DIALOG,
            directory=os.path.dirname(ruta_doc),
            save_filename=f"anotado_{nombre_base}.pdf"
        )

        if not resultado:
            return {"exito": False}

        ruta_destino = resultado if isinstance(resultado, str) else resultado[0]

        try:
            anotaciones = json.loads(datos_json)
            pdf_bytes = base64.b64decode(datos_b64)
            doc_fitz = fitz.open(stream=pdf_bytes, filetype="pdf")

            # Resaltados y tachados con soporte para colores personalizados
            for item in anotaciones.get("resaltados", []):
                pag = doc_fitz[item["pageNum"] - 1]
                rect = fitz.Rect(item["x"], item["y"], item["x"] + item["w"], item["y"] + item["h"])
                if item.get("tipo") == "tachado":
                    annot = pag.add_strikeout_annot(rect)
                else:
                    annot = pag.add_highlight_annot(rect)
                    col_hex = item.get("color", "#facc15").lstrip("#")
                    col = tuple(int(col_hex[i:i+2], 16) / 255.0 for i in (0, 2, 4))
                    annot.set_colors(stroke=col)
                annot.update()

            # Trazos del lápiz marcador libre
            for tr in anotaciones.get("trazos", []):
                puntos = tr.get("puntos", [])
                if len(puntos) > 1:
                    pag = doc_fitz[tr["pageNum"] - 1]
                    fitz_pts = [fitz.Point(pt["x"], pt["y"]) for pt in puntos]
                    annot = pag.add_ink_annot([fitz_pts])
                    col_hex = tr.get("color", "#facc15").lstrip("#")
                    col = tuple(int(col_hex[i:i+2], 16) / 255.0 for i in (0, 2, 4))
                    annot.set_colors(stroke=col)
                    annot.set_border(width=tr.get("grosor", 18))
                    annot.update()

            # Textos insertados
            for t in anotaciones.get("textos", []):
                texto_limpio = t.get("texto", "").strip()
                if not texto_limpio:
                    continue
                pag = doc_fitz[t["pageNum"] - 1]
                punto = fitz.Point(t["x"], t["y"] + t.get("fontSize", 14))
                col_hex = t.get("color", "#ea580c").lstrip("#")
                col = tuple(int(col_hex[i:i+2], 16) / 255.0 for i in (0, 2, 4))
                pag.insert_text(punto, texto_limpio, fontsize=t.get("fontSize", 14), color=col, fontname="helv")

            # Comentarios
            for com in anotaciones.get("comentarios", []):
                pag = doc_fitz[com["pageNum"] - 1]
                punto = fitz.Point(com["relX"], com["relY"])
                annot = pag.add_text_annot(punto, f"[{com.get('hora', 'Hoy')}]\n{com['texto']}")
                annot.update()

            # Guardar marcadores personalizados dentro del PDF (TOC)
            marcadores = anotaciones.get("marcadores", [])
            if marcadores:
                toc = [[1, m["titulo"], m["pageNum"]] for m in marcadores]
                doc_fitz.set_toc(toc)

            doc_fitz.save(ruta_destino)
            return {"exito": True, "ruta": ruta_destino}
        except Exception as e:
            return {"exito": False, "error": str(e)}


def iniciar_servidor_local(puerto, directorio):
    class Handler(SimpleHTTPRequestHandler):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, directory=directorio, **kwargs)
        def log_message(self, format, *args):
            pass  # Silenciar logs en terminal
        def end_headers(self):
            self.send_header("Access-Control-Allow-Origin", "*")
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
            self.send_header("Access-Control-Allow-Headers", "*")
            super().end_headers()
        def do_OPTIONS(self):
            self.send_response(200)
            self.end_headers()
        def do_POST(self):
            if self.path == "/api/convert":
                try:
                    content_length = int(self.headers.get("Content-Length", 0))
                    body = self.rfile.read(content_length)
                    data = json.loads(body.decode("utf-8"))
                    nombre = data.get("nombre", "doc.pptx")
                    b64 = data.get("data", "")
                    temp_in = os.path.join(tempfile.gettempdir(), f"input_api_{nombre}")
                    with open(temp_in, "wb") as f:
                        f.write(base64.b64decode(b64))
                    pdf_bytes = convertir_a_pdf(temp_in)
                    try: os.remove(temp_in)
                    except Exception: pass
                    if pdf_bytes:
                        self.send_response(200)
                        self.send_header("Content-Type", "application/json")
                        self.end_headers()
                        res = json.dumps({"exito": True, "data": base64.b64encode(pdf_bytes).decode("utf-8")})
                        self.wfile.write(res.encode("utf-8"))
                        return
                except Exception:
                    pass
                self.send_response(500)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(b'{"exito": false}')
                return
            super().do_POST()

    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("127.0.0.1", puerto), Handler) as httpd:
        httpd.serve_forever()

def on_closing(window):
    def consultar():
        window.evaluate_js("iniciarFlujoCierreApp()")
    
    threading.Thread(target=consultar, daemon=True).start()
    return False  # Cancela el cierre abrupto para permitir que el modal responda

def main():
    api = PDFViewerAPI()
    dir_actual = os.path.dirname(os.path.abspath(__file__))
    puerto = 8765

    hilo_server = threading.Thread(
        target=iniciar_servidor_local, 
        args=(puerto, dir_actual), 
        daemon=True
    )
    hilo_server.start()

    window = webview.create_window(
        title="Lector Documental Moderno — Acelerado por GPU",
        url=f"http://127.0.0.1:{puerto}/index.html",
        js_api=api,
        width=1440,
        height=900,
        maximized=True,
        background_color="#0e0e11",
        easy_drag=False
    )
    api.set_window(window)
    window.events.closing += lambda: on_closing(window)
    webview.start(gui="edgechromium", debug=False)

if __name__ == "__main__":
    main()